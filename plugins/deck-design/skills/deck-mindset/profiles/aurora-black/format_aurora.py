"""Deck format helpers — Slidecast "Aurora Tech" vocabulary applied to deck-mindset rules.

Color vocabulary:  TEXT #EAF0FF · ACCENT #FF40FF (magenta, the one focal color) · TEAL (structure labels) · MUTED #D6DCEA · BOX_LINE #C9D1E3
Chrome: template provides background + AWS logo + copyright; code adds only the bottom cite (title + https). No bars, no page numbers.
Type scale (max 3 sizes per slide): TITLE 30 / MID 18 / BODY 15 (+ CITE 10 caption exception)
Slide: 13.333 x 7.5 in  (1px of the 1920-wide Slidecast == 0.5pt)
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
import copy, os

# ── vocabulary ────────────────────────────────────────────────────────
BG     = RGBColor(0x06, 0x0B, 0x1A)
TEXT   = RGBColor(0xEA, 0xF0, 0xFF)
ACCENT = RGBColor(0xFF, 0x40, 0xFF)   # magenta (DI deck accent) — the one focal color, user-chosen
TEAL   = RGBColor(0x2B, 0xD9, 0xC7)   # secondary: lane/structure labels only
MUTED  = RGBColor(0xD6, 0xDC, 0xEA)   # captions + small-tier descriptions — light gray (≈ white 85%)
GRAD   = ["2BD9C7", "4DA3FF", "9B5CFF"]  # teal → blue → purple (hero/gradient bar)
SURFACE_LINE = RGBColor(0x5B, 0x64, 0x78)  # rules / lifelines (≈ white 35% on BG)
BOX_LINE     = RGBColor(0xC9, 0xD1, 0xE3)  # card / box borders — light, near white
SURFACE_FILL = RGBColor(0x14, 0x1B, 0x30)  # glass card fill (≈ white 8% on BG)

TITLE, MID, BODY, CITE = 30, 18, 15, 10   # CITE 12→10 (2026-09-05): keep every reference, smaller type
# RENDER_CHECK=1 → use Noto for latin too, so LibreOffice QA renders match PowerPoint spacing
RENDER_CHECK = os.environ.get("RENDER_CHECK") == "1"
FONT_LATIN = "Noto Sans CJK KR" if RENDER_CHECK else "Amazon Ember"
FONT_EA    = "Noto Sans CJK KR"
FONT_MONO  = "Noto Sans Mono CJK KR" if RENDER_CHECK else "JetBrains Mono"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN_X = Inches(0.83)   # 120px
ASSETS = os.environ.get("AURORA_ASSETS", os.path.join(os.path.dirname(__file__), "assets"))


BASE = os.path.join(ASSETS, "base-black.pptx")   # DI deck master: black gradient bg + AWS logo + copyright in layouts


def new_presentation():
    """Open the slim black template (master + layouts only). Background, logo and copyright come from the layout."""
    return Presentation(BASE)


def layout(prs, name="1_Blank"):
    for l in prs.slide_layouts:
        if l.name == name:
            return l
    raise KeyError(name)


# ── text helpers ──────────────────────────────────────────────────────
def _set_fonts(run, size, bold=False, color=TEXT, mono=False, gradient=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.name = FONT_MONO if mono else FONT_LATIN
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT_EA)
    if gradient:
        # gradient text fill (PowerPoint); LibreOffice may fall back to first stop
        for old in rPr.findall(qn("a:solidFill")) + rPr.findall(qn("a:gradFill")):
            rPr.remove(old)
        gf = etree.Element(qn("a:gradFill"))
        gs = etree.SubElement(gf, qn("a:gsLst"))
        for pos, hexv in zip((0, 45000, 100000), GRAD):
            g = etree.SubElement(gs, qn("a:gs")); g.set("pos", str(pos))
            c = etree.SubElement(g, qn("a:srgbClr")); c.set("val", hexv)
        lin = etree.SubElement(gf, qn("a:lin")); lin.set("ang", "0"); lin.set("scaled", "0")
        rPr.insert(0, gf)
    else:
        f.color.rgb = color


def add_text(slide, left, top, width, height, lines, size=BODY, bold=False, color=TEXT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, mono=False, gradient=False,
             line_spacing=1.15, space_after=0):
    """lines: str | list[str | list[(text, opts)]]  — opts: size,bold,color,mono,gradient"""
    left, top, width, height = int(left), int(top), int(width), int(height)
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(lines, str):
        lines = [lines]
    # expand embedded newlines in plain strings into separate paragraphs
    expanded = []
    for line in lines:
        if isinstance(line, str) and "\n" in line:
            expanded.extend(line.split("\n"))
        else:
            expanded.append(line)
    lines = expanded
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        segs = line if isinstance(line, list) else [(line, {})]
        for text, opts in segs:
            r = p.add_run(); r.text = text
            _set_fonts(r, opts.get("size", size), opts.get("bold", bold), opts.get("color", color),
                       opts.get("mono", mono), opts.get("gradient", gradient))
    return tb


# ── chrome: background, gradient bar, footer, cite, page number ───────
def add_chrome(slide, cite=None, cover=False):
    """Per-slide chrome on top of the template layout: optional cite only. No bar, no page numbers (user notes 2026-09-03)."""
    if cover:
        return
    if cite:
        # cite: str | list[str] — every public reference carries title + https URL (one per line)
        lines = [cite] if isinstance(cite, str) else list(cite)
        h = Inches(0.19) * len(lines)
        CITE_LEFT, CITE_BOTTOM = MARGIN_X, Inches(6.98)   # full width, bottom edge just above the template copyright (y ≈ 7.05in)
        add_text(slide, CITE_LEFT, CITE_BOTTOM - h, SLIDE_W - CITE_LEFT - MARGIN_X, h, lines,
                 size=CITE, color=MUTED, align=PP_ALIGN.RIGHT, line_spacing=1.05)


TITLE_MAX_CHARS = 34   # ≈ one line at 30pt bold on 11.67in (Korean ≈ 0.34in/char)


TITLE_TOP = Inches(0.32)   # user note: title close to the top to free vertical space


def add_title(slide, text, top=TITLE_TOP):
    if len(text) > TITLE_MAX_CHARS:
        print(f"[title too long: {len(text)} chars] {text}")
    return add_text(slide, MARGIN_X, top, SLIDE_W - 2 * MARGIN_X, Inches(0.7), text, size=TITLE, bold=True)


# ── shapes: hairline, glass card, box, arrow ──────────────────────────
def _rule(slide, left, top, width, height, color):
    """Thin filled rectangle used as a rule (avoids zero-size connector warnings in Phase-1 QA)."""
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(left), int(top), int(width), int(height))
    r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background()
    return r


def add_hairline(slide, left, top, width, color=SURFACE_LINE, weight=0.75):
    return _rule(slide, left, top, width, Pt(weight), color)


def add_vline(slide, left, top, height, color=SURFACE_LINE, weight=0.75, dash=False):
    if not dash:
        return _rule(slide, left, top, Pt(weight), height, color)
    # dashed lifeline = stack of short rules
    seg, gap = Inches(0.12), Inches(0.08)
    y = top
    while y < top + height:
        _rule(slide, left, y, Pt(weight), min(seg, top + height - y), color)
        y += seg + gap


def add_box(slide, left, top, width, height, text_lines, size=BODY, bold=False, color=TEXT,
            fill=SURFACE_FILL, line=BOX_LINE, align=PP_ALIGN.CENTER, radius=True):
    left, top, width, height = int(left), int(top), int(width), int(height)
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 left, top, width, height)
    if radius:
        shp.adjustments[0] = 0.12
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(0.75)
    tf = shp.text_frame; tf.word_wrap = True
    # inner padding so text never touches the border (rule: ≥0.12in sides, ≥0.08in top/bottom)
    tf.margin_left = tf.margin_right = Inches(0.12); tf.margin_top = tf.margin_bottom = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if isinstance(text_lines, str):
        text_lines = [text_lines]
    for i, line in enumerate(text_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        segs = line if isinstance(line, list) else [(line, {})]
        for t, o in segs:
            r = p.add_run(); r.text = t
            _set_fonts(r, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("mono", False))
    return shp


def add_arrow(slide, x1, y1, x2, y2, color=ACCENT, weight=1.25, head=True):
    # 1 EMU offset keeps the connector non-degenerate for the Phase-1 zero-size check
    x1, y1, x2, y2 = int(x1), int(y1), int(x2), int(y2)
    if y1 == y2: y2 = y2 + 1
    if x1 == x2: x2 = x2 + 1
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color; c.line.width = Pt(weight)
    if head:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd")); tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    return c


def set_notes(slide, text):
    """Speaker notes (발표자 노트) — plain text."""
    slide.notes_slide.notes_text_frame.text = text
