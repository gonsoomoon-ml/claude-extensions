"""light profile — deck-mindset rules on a plain white background (python-pptx).

No template file is needed: new_presentation() starts from python-pptx's default deck, resized to
13.333 x 7.5 in. Set LIGHT_TEMPLATE=/path/to/your-light-template.pptx to build on your own template instead
(its sample slides are dropped; masters, layouts and theme are kept — check LAYOUT_CONTENT / LAYOUT_STATEMENT).

- Colors (fixed roles): TEXT body · MUTED cites and secondary marks only · ACCENT one focal subject per slide,
  plus two semantic colors (BLUE, ORANGE) — give each one meaning per deck and keep it on every slide.
- Type: HEADLINE 40 · TITLE 30 · TAKEAWAY 24 · MID 18 · BODY 15 (+ CITE 10 for citations), three per slide.
  Slide code uses role names (ROLE), never literal sizes. check_slides() runs before save and stops the build
  on a violation — it does not warn and continue.
- Fonts: Latin in Amazon Ember; Korean by the machine that presents — Noto Sans CJK KR by default,
  DECK_FONT_KR="Apple SD Gothic Neo" for a Mac. RENDER_CHECK=1 renders both in Noto Sans CJK KR (QA on Linux).
- No chrome: no logo, footer or page number. The cite line is the only furniture. No speaker notes, no
  animations — the presenter animates by hand, one named group per reveal step.
"""
import os
import re
import subprocess
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_LINE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from lxml import etree

# ── colors (contrast on white: TEXT 17.0 · MUTED 5.9 · ACCENT 5.0 · BLUE_TEXT 4.8 · ORANGE_TEXT 5.2) ──
TEXT        = RGBColor(0x16, 0x1D, 0x26)
MUTED       = RGBColor(0x5B, 0x65, 0x73)   # cites and secondary marks only — never audience prose
ACCENT      = RGBColor(0x8A, 0x3F, 0xFC)   # the one focal color
ACCENT_TINT = RGBColor(0xF3, 0xEC, 0xFF)
BLUE        = RGBColor(0x1A, 0x8C, 0xFF)   # semantic color 1 — borders and fills (3.4:1, not for text)
BLUE_TEXT   = RGBColor(0x09, 0x72, 0xD3)   # semantic color 1 — text
BLUE_TINT   = RGBColor(0xE8, 0xF3, 0xFF)
ORANGE      = RGBColor(0xE8, 0x54, 0x1E)   # semantic color 2 — borders and fills (3.7:1, not for text)
ORANGE_TEXT = RGBColor(0xC2, 0x41, 0x0C)   # semantic color 2 — text
ORANGE_TINT = RGBColor(0xFF, 0xEE, 0xE7)
BOX_LINE    = RGBColor(0xC5, 0xCC, 0xD6)
CARD_FILL   = RGBColor(0xF7, 0xF7, 0xFA)
TILE_FILL   = RGBColor(0xEE, 0xF1, 0xF5)   # map tiles — visible on white without a border
RULE        = RGBColor(0xD5, 0xDA, 0xE1)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)

HEADLINE, TITLE, TAKEAWAY, MID, BODY, CITE = 40, 30, 24, 18, 15, 10
SCALE = {HEADLINE, TITLE, TAKEAWAY, MID, BODY, CITE}   # check_slides() rejects any other size (15pt floor included)

# role → (size, bold, color). Slide code names a role; add a role when a new kind of text is approved.
# 18 and 15 differ by only 1.2x, so on a slide that has 15 every 18 must be bold or the accent (checked).
ROLE = {
    "headline":      (HEADLINE, True,  TEXT),    # the sentence of a statement slide
    "title":         (TITLE,    True,  TEXT),    # content-slide title, in the title placeholder
    "hero_question": (TITLE,    True,  TEXT),    # the question line on a statement slide
    "badge":         (TITLE,    True,  WHITE),   # A-D letter in a filled circle
    "takeaway":      (TAKEAWAY, True,  TEXT),    # one-line conclusion — no 18 on that slide
    "question":      (MID,      True,  TEXT),    # a question line on a content slide
    "bridge":        (MID,      True,  TEXT),    # the line that hands over to the next slide
    "callout":       (MID,      True,  ACCENT),  # names the slide's focus in a few words
    "card":          (MID,      False, TEXT),    # card text — regular, so not on a slide that has 15
    "lead":          (MID,      False, TEXT),    # explanation under a headline — regular, so not with 15
    "label":         (BODY,     True,  TEXT),    # column head, item name
    "body":          (BODY,     False, TEXT),    # description, condition
    "cite":          (CITE,     False, MUTED),   # citation line only
}

RENDER_CHECK = os.environ.get("RENDER_CHECK") == "1"
FONT_LATIN = "Noto Sans CJK KR" if RENDER_CHECK else "Amazon Ember"
FONT_EA = "Noto Sans CJK KR" if RENDER_CHECK else os.environ.get("DECK_FONT_KR", "Noto Sans CJK KR")

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
X0 = Inches(0.67)                 # left edge of the title and the content
CONTENT_W = Inches(12.0)
TITLE_TOP = Inches(0.15)          # title close to the top — the same spot on every content slide
TITLE_H = Inches(0.65)
CONTENT_TOP = Inches(1.0)         # below the title

LAYOUT_CONTENT = "Title Only"     # content slides — a real title placeholder (outline view, screen readers)
LAYOUT_STATEMENT = "Blank"        # statement / question slides


def new_presentation():
    """A 13.333 x 7.5 in deck: python-pptx's default (white), or LIGHT_TEMPLATE with its sample slides dropped."""
    path = os.environ.get("LIGHT_TEMPLATE")
    prs = Presentation(path) if path else Presentation()
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)
    if not path:
        prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    elif abs(prs.slide_width - SLIDE_W) > Inches(0.01) or abs(prs.slide_height - SLIDE_H) > Inches(0.01):
        # every position here assumes 16:9 — a 4:3 template would push shapes off the slide, silently
        raise SystemExit(f"LIGHT_TEMPLATE is {prs.slide_width / 914400:.2f} x {prs.slide_height / 914400:.2f} in — "
                         "this profile lays out 13.333 x 7.5 in (16:9)")
    return prs


def layout(prs, name, master=0):
    for l in prs.slide_masters[master].slide_layouts:
        if l.name == name:
            return l
    raise KeyError(name)


def content_slide(prs, title):
    s = prs.slides.add_slide(layout(prs, LAYOUT_CONTENT))
    set_title(s, title)
    return s


def statement_slide(prs):
    return prs.slides.add_slide(layout(prs, LAYOUT_STATEMENT))


# ── text ──────────────────────────────────────────────────────────────
def _font(run, size, bold=False, color=TEXT):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.name = FONT_LATIN
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):           # Korean face in ea/cs on every run
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT_EA)


def _fill_tf(tf, lines, size, bold, color, align, line_spacing=1.1, space_after=0):
    if isinstance(lines, str):
        lines = lines.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        segs = line if isinstance(line, list) else [(line, {})]
        for text, o in segs:
            r = p.add_run()
            r.text = text
            s, b, c = ROLE[o["role"]] if "role" in o else (size, bold, color)
            _font(r, o.get("size", s), o.get("bold", b), o.get("color", c))


def add_text(container, left, top, width, height, lines, role="body", align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.1, space_after=0):
    """lines: str | list[str | list[(text, opts)]] — opts may set role, bold or color for one run."""
    size, bold, color = ROLE[role]
    tb = container.shapes.add_textbox(int(left), int(top), int(width), int(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    _fill_tf(tf, lines, size, bold, color, align, line_spacing, space_after)
    return tb


def set_title(slide, text):
    """Fill the title placeholder at TITLE_TOP. All four of left/top/width/height are written: writing only
    `top` on an inherited placeholder leaves no size and the title jumps to the left edge. Autofit is off, so
    a long title does not shrink — check_slides() catches it and the sentence gets shorter instead."""
    ph = slide.shapes.title
    ph.left, ph.top, ph.width, ph.height = X0, TITLE_TOP, CONTENT_W, TITLE_H
    tf = ph.text_frame
    tf.clear()
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    size, bold, color = ROLE["title"]
    _fill_tf(tf, [text], size, bold, color, PP_ALIGN.LEFT)
    return ph


def add_cite(slide, lines):
    """Citations: title + https, one per line, 10pt MUTED, full width, bottom right. Shown from the start."""
    lines = [lines] if isinstance(lines, str) else list(lines)
    h = Inches(0.19) * len(lines)
    tb = add_text(slide, X0, Inches(7.15) - h, CONTENT_W, h, lines, role="cite", align=PP_ALIGN.RIGHT,
                  line_spacing=1.0)
    tb.name = "start · cite"
    return tb


# ── shapes ────────────────────────────────────────────────────────────
def _flat(shp):
    """No shadow. The default theme's effect styles carry outer shadows; an empty effectLst alone is honored
    by PowerPoint but not by LibreOffice QA renders, so the shape's style reference is set to "no effect" too."""
    shp.shadow.inherit = False
    style = shp._element.find(qn("p:style"))
    if style is not None and style.find(qn("a:effectRef")) is not None:
        style.find(qn("a:effectRef")).set("idx", "0")


def add_group(slide, name):
    """One reveal step. The name ("click1 · …", "start · …") shows in PowerPoint's Selection Pane, where the
    presenter gives the group one animation. Pass the group instead of the slide to draw inside it."""
    g = slide.shapes.add_group_shape()
    g.name = name
    return g


def add_rule(container, left, top, width, color=RULE, weight=0.75):
    """A thin filled rectangle as a rule — a zero-height connector is flagged by some validators."""
    r = container.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(left), int(top), int(width), int(Pt(weight)))
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    _flat(r)
    return r


def add_box(container, left, top, width, height, lines=None, role="body", fill=CARD_FILL, line=BOX_LINE,
            line_w=0.75, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, radius=0.08, pad_x=0.16, pad_y=0.12,
            line_spacing=1.1, space_after=0, dash=False):
    """Card or box with optional text. fill=None / line=None for no fill / no border."""
    size, bold, color = ROLE[role]
    shp = container.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                     int(left), int(top), int(width), int(height))
    if radius:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    _flat(shp)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(pad_x)
    tf.margin_top = tf.margin_bottom = Inches(pad_y)
    tf.vertical_anchor = anchor
    if lines:
        _fill_tf(tf, lines, size, bold, color, align, line_spacing, space_after)
    return shp


def add_frame(container, card, color=ACCENT, weight=2.25, radius=0.08):
    """Highlight a card as a later reveal step: a border at the card's exact geometry, drawn on top. The wider
    line covers the card's own border — a frame offset outward shows a double border."""
    return add_box(container, card.left, card.top, card.width, card.height, fill=None, line=color,
                   line_w=weight, radius=radius)


def add_veil(container, card, alpha=40, radius=0.08):
    """Dim a card as a later reveal step: a white shape at alpha % opacity on top. An animation can reveal a
    shape but cannot restyle an existing one, so a change of state is always a new shape."""
    pad = Inches(0.02)
    v = add_box(container, card.left - pad, card.top - pad, card.width + 2 * pad, card.height + 2 * pad,
                fill=WHITE, line=None, radius=radius)
    clr = v.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    etree.SubElement(clr, qn("a:alpha")).set("val", str(alpha * 1000))
    return v


def add_badge(container, left, top, letter, size=0.55, fill=TEXT):
    """Letter badge for an option card — white letter in a filled circle (role 'badge')."""
    b = container.shapes.add_shape(MSO_SHAPE.OVAL, int(left), int(top), int(Inches(size)), int(Inches(size)))
    b.fill.solid()
    b.fill.fore_color.rgb = fill
    b.line.fill.background()
    _flat(b)
    tf = b.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    size_pt, bold, color = ROLE["badge"]
    _fill_tf(tf, [letter], size_pt, bold, color, PP_ALIGN.CENTER, line_spacing=1.0)
    return b


def add_marker(container, left, top, size=0.13, color=ACCENT, rotation=90):
    """Small triangle pointer — rotation 90 = right (▶), 0 = up (▲). A shape, not a glyph, so its size and
    color follow the tokens on any machine."""
    m = container.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, int(left), int(top),
                                   int(Inches(size)), int(Inches(size)))
    m.rotation = rotation
    m.fill.solid()
    m.fill.fore_color.rgb = color
    m.line.fill.background()
    _flat(m)
    return m


def add_block_arrow(container, left, top, width, height, direction="right", color=ACCENT):
    """Block arrow (a thick filled shape) — the slide's one core action. Correspondence and flow: add_arrow."""
    kind = {"left": MSO_SHAPE.LEFT_ARROW, "right": MSO_SHAPE.RIGHT_ARROW,
            "up": MSO_SHAPE.UP_ARROW, "down": MSO_SHAPE.DOWN_ARROW}[direction]
    a = container.shapes.add_shape(kind, int(left), int(top), int(width), int(height))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    _flat(a)
    return a


def _connector(container, x1, y1, x2, y2, color, weight):
    x1, y1, x2, y2 = int(x1), int(y1), int(x2), int(y2)
    if y1 == y2:      # 1 EMU offset — a zero-size connector is flagged by some validators
        y2 += 1
    if x1 == x2:
        x2 += 1
    c = container.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    return c


def add_line(container, x1, y1, x2, y2, color=MUTED, weight=1.5, dash=False):
    """Line without a head (brackets, scale bars)."""
    c = _connector(container, x1, y1, x2, y2, color, weight)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    return c


def add_arrow(container, x1, y1, x2, y2, color=MUTED, weight=1.25, both=False):
    """Line arrow — a correspondence or flow; light enough to repeat."""
    c = _connector(container, x1, y1, x2, y2, color, weight)
    ln = c.line._get_or_add_ln()
    if both:  # schema order: headEnd before tailEnd
        head = etree.SubElement(ln, qn("a:headEnd"))
        head.set("type", "triangle"); head.set("w", "med"); head.set("len", "med")
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    return c


def drop_empty_placeholders(slide):
    """Remove placeholders left empty (an unused title, date, footer, number)."""
    for ph in list(slide.placeholders):
        if not (ph.has_text_frame and ph.text_frame.text.strip()):
            ph._element.getparent().remove(ph._element)


# ── build check ───────────────────────────────────────────────────────
# The width estimate is calibrated to Noto Sans CJK KR and LibreOffice QA renders (2026-10, one deck):
# Korean wraps at spaces, a line at spacing 1.0 is ~1.23em tall, Amazon Ember Latin may run up to 10% wider
# than Noto. Re-calibrate against your own renders before trusting it with other fonts or renderers.
_fonts = {}   # bold → ImageFont at 1000 units | False (not found)


def _measure_font(bold=True):
    """Noto Sans CJK KR via fc-match + Pillow; None if unavailable (falls back to _em)."""
    if bold not in _fonts:
        _fonts[bold] = False
        try:
            from PIL import ImageFont
            pattern = "Noto Sans CJK KR:bold" if bold else "Noto Sans CJK KR:regular"
            path = subprocess.run(["fc-match", "-f", "%{file}", pattern],
                                  capture_output=True, text=True).stdout.strip()
            for i in range(8):  # find the KR face inside the .ttc
                f = ImageFont.truetype(path, 1000, index=i)
                if "KR" in f.getname()[0]:
                    _fonts[bold] = f
                    break
        except Exception:
            pass
    return _fonts[bold] or None


def _em(ch):
    """Character width in em when the font is missing — matched to Noto Sans CJK KR Bold."""
    if "가" <= ch <= "힣" or "　" <= ch <= "鿿":
        return 0.92
    if ch == " ":
        return 0.23
    if ch.isdigit():
        return 0.59
    if ch.isascii() and ch.isalpha():
        return 0.66 if ch.isupper() else 0.60
    return 0.60


def width_in(text, size, bold=True):
    """Width of a run of text in inches, +up to 10% for the Latin share (Amazon Ember vs Noto)."""
    f = _measure_font(bold)
    em = f.getlength(text) / 1000 if f else sum(_em(c) for c in text)
    latin = sum(c.isascii() and c.isalpha() for c in text) / max(len(text), 1)
    return em * size / 72 * (1 + 0.10 * latin)


_TOKEN = re.compile(r"\S+|\s+")   # wrap word by word (measured: Korean also wraps at spaces in QA renders)
LINE_PITCH = 1.25                 # line height in em at spacing 1.0 — measured 1.23, rounded up


def _paragraph_lines(p, room):
    """Lines a paragraph takes at `room` inches, and its largest size."""
    lines, x, biggest = 1, 0.0, 0
    for el in p._p.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "br":
            lines, x = lines + 1, 0.0
            continue
        if tag not in ("r", "fld"):
            continue
        rPr = el.find(qn("a:rPr"))
        size = int(rPr.get("sz")) / 100 if rPr is not None and rPr.get("sz") else BODY
        bold = rPr is not None and rPr.get("b") == "1"
        biggest = max(biggest, size)
        t = el.find(qn("a:t"))
        for tok in _TOKEN.findall(t.text or "") if t is not None else []:
            w = width_in(tok, size, bold)
            if x > 0 and x + w > room and not tok.isspace():
                lines, x = lines + 1, 0.0
            if not (x == 0 and tok.isspace()):
                x += w
    return lines, biggest


def _box_overflow(sh):
    """(needed, available) height in inches if the text overflows its box, else None."""
    tf = sh.text_frame
    room = (sh.width - tf.margin_left - tf.margin_right) / 914400 if tf.word_wrap is not False else 1e9
    avail = (sh.height - tf.margin_top - tf.margin_bottom) / 914400
    need, paras = 0.0, tf.paragraphs
    for i, p in enumerate(paras):
        n, size = _paragraph_lines(p, room)
        spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        need += n * (size or BODY) * LINE_PITCH * spacing / 72
        if i < len(paras) - 1 and p.space_after:
            need += p.space_after / 914400
    return (need, avail) if need > avail + 0.03 else None


def _walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(sh.shapes)
        else:
            yield sh


def _rgb(run):
    c = run.font.color
    return c.rgb if c.type == MSO_COLOR_TYPE.RGB else None


def check_slides(prs):
    """Stop the build on a violation — call it before prs.save(); nothing is saved if it fails.
    1 sizes only from SCALE  2 three sizes per slide (cite excluded)  3 on a slide with 15, an 18 is bold or
    the accent  4 the title fits one line — shorten the sentence, never the font  5 no text overflows its box."""
    problems = []
    for n, slide in enumerate(prs.slides, 1):
        runs = [r for sh in _walk(slide.shapes) if sh.has_text_frame
                for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()]
        sizes = set()
        for r in runs:
            if r.font.size is None:
                problems.append(f"slide {n}: no size on '{r.text[:24]}'")
                continue
            sizes.add(r.font.size.pt)
            if r.font.size.pt not in SCALE:
                problems.append(f"slide {n}: {r.font.size.pt:g}pt is not in the scale — '{r.text[:24]}'")
        used = sizes - {CITE}
        if len(used) > 3:
            problems.append(f"slide {n}: {len(used)} sizes {sorted(used, reverse=True)} — three at most")
        if BODY in used:
            for r in runs:
                if r.font.size is not None and r.font.size.pt == MID and not r.font.bold and _rgb(r) != ACCENT:
                    problems.append(f"slide {n}: an 18 beside a 15 must be bold or the accent — '{r.text[:24]}'")
        t = slide.shapes.title
        for sh in _walk(slide.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip() and not (t is not None and sh.shape_id == t.shape_id):
                over = _box_overflow(sh)
                if over:
                    problems.append(f"slide {n}: text overflows ({over[0]:.2f}in > box {over[1]:.2f}in) "
                                    f"'{sh.text_frame.text[:24]}' — cut the text or grow the box")
        if t is not None and t.text_frame.text.strip():
            text = t.text_frame.text
            # 3% slack — a full one-line title breaks in LibreOffice QA renders (extra gaps around Latin)
            room = (t.width - t.text_frame.margin_left - t.text_frame.margin_right) / 914400 * 0.97
            w = width_in(text, TITLE, True)
            if "\n" in text or "\v" in text or w > room:
                problems.append(f"slide {n}: title longer than one line ({w:.2f}in > {room:.2f}in) — shorten it")
    if problems:
        print("build check failed — nothing saved:", *problems, sep="\n  ", file=sys.stderr)
        raise SystemExit(1)


def count_text(prs):
    """Print characters per slide (spaces, cites and Latin words excluded — mostly product names).
    Print only: the budget is kept when the mock is approved."""
    counts = []
    for slide in prs.slides:
        n = 0
        for sh in _walk(slide.shapes):
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size is None or r.font.size.pt != CITE:
                            n += len(re.sub(r"\s", "", re.sub(r"[A-Za-z][A-Za-z0-9.\-]*", "", r.text)))
        counts.append(n)
    print("characters per slide:", " · ".join(f"{i} {n}" for i, n in enumerate(counts, 1)), flush=True)
