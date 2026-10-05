"""Three example slides on the light profile — statement + option cards, block arrow into slots,
answers revealed with frames, veils and line arrows. Every reveal step is a named group.

    RENDER_CHECK=1 python3 example_build.py          # QA build (Noto for both faces) → out/example_light.pptx
    python3 ../../../deck-build/scripts/qa.py out/example_light.pptx
"""
import os
from pptx.util import Inches as I
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from format_light import (
    new_presentation, content_slide, statement_slide, add_text, add_cite, add_box, add_group, add_frame,
    add_veil, add_badge, add_block_arrow, add_arrow, drop_empty_placeholders, check_slides,
    count_text, X0, CONTENT_W, ACCENT, MUTED, WHITE, BLUE, BLUE_TEXT, BLUE_TINT, ORANGE, ORANGE_TEXT,
    ORANGE_TINT,
)

LEFT, RIGHT = 0.67, 12.67
OPTIONS = ["서버를 직접 구성한다", "컨테이너로 묶어 올린다", "관리형 서비스에 맡긴다", "다음 분기로 미룬다"]
CARD_W, CARD_GAP, BADGE, PAD = 2.85, 0.2, 0.55, 0.16       # 4 x 2.85 + 3 x 0.2 = 12.0 — the content width


def card_x(i):
    return LEFT + i * (CARD_W + CARD_GAP)


def option_cards(container, y, h, role):
    """Four equal option cards with A-D badges; returns the cards (for frames and veils later)."""
    cards = []
    for i, text in enumerate(OPTIONS):
        c = add_box(container, I(card_x(i)), I(y), I(CARD_W), I(h), text, role=role, fill=WHITE, line=MUTED,
                    line_w=1.0)
        c.text_frame.margin_top = I(PAD + BADGE + 0.1)          # text starts below the badge
        c.name = f"card {'ABCD'[i]}"
        cards.append(c)
    for i in range(4):                                         # badges last, so they sit on top
        add_badge(container, I(card_x(i) + PAD), I(y + PAD), "ABCD"[i], size=BADGE).name = f"badge {'ABCD'[i]}"
    return cards


prs = new_presentation()

# ── 1 · statement + options ── start = scene → click1 = question → click2 = options
s = statement_slide(prs)
add_text(s, X0, I(1.1), CONTENT_W, I(1.5),
         ["새 서비스를 다음 달까지", [("운영에 올려야", {"color": ACCENT}), (" 한다면", {})]],
         role="headline", line_spacing=1.0).name = "start · scene"
add_text(s, X0, I(2.95), CONTENT_W, I(0.6), "어떻게 배포하시겠습니까?", role="hero_question").name = "click1 · question"
option_cards(add_group(s, "click2 · options"), 3.95, 1.6, "card")
drop_empty_placeholders(s)

# ── 2 · block arrow = the slide's one core action (plug the pieces in)
# start = platform + pieces → click1 = block arrows → click2 = bridge line
s = content_slide(prs, "직접 쓴 세 조각을 플랫폼에 끼운다")
PANEL_R, PIECE_X, ROWS = 7.75, 8.6, [2.05, 3.05, 4.05]
g = add_group(s, "start · platform")
add_box(g, I(LEFT), I(1.2), I(PANEL_R - LEFT), I(3.85), fill=BLUE_TINT, line=BLUE, line_w=1.25).name = "platform"
add_text(g, I(LEFT + 0.25), I(1.4), I(4.5), I(0.3),
         [[("관리형 플랫폼", {"role": "label", "color": BLUE_TEXT}), ("  가져다 쓴다", {"role": "body"})]],
         role="body").name = "platform head"
for k, chip in enumerate(["실행", "확장", "로그"]):
    add_box(g, I(LEFT + 0.25), I(ROWS[k]), I(2.2), I(0.7), chip, role="body", fill=WHITE, line=None,
            anchor=MSO_ANCHOR.MIDDLE).name = f"chip {chip}"
for k in range(3):
    add_box(g, I(5.55), I(ROWS[k]), I(1.95), I(0.7), "끼울 자리", role="body", fill=WHITE, line=BLUE, dash=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE).name = f"slot {k + 1}"
g = add_group(s, "start · pieces")
add_text(g, I(PIECE_X), I(1.4), I(RIGHT - PIECE_X), I(0.3), [[("직접 정의한다", {"role": "label", "color": ORANGE_TEXT})]],
         role="label").name = "pieces head"
for k, (name, detail) in enumerate([("설정 파일", "환경마다 다른 값"), ("배포 스크립트", "올리는 순서"),
                                    ("경보 규칙", "언제 깨울지")]):
    add_box(g, I(PIECE_X), I(ROWS[k]), I(RIGHT - PIECE_X), I(0.7),
            [[(name, {"role": "label", "color": ORANGE_TEXT}), ("  " + detail, {"role": "body"})]],
            fill=ORANGE_TINT, line=ORANGE, anchor=MSO_ANCHOR.MIDDLE).name = f"piece {name}"
g = add_group(s, "click1 · plug in")
for k in range(3):
    add_block_arrow(g, I(7.8), I(ROWS[k] + 0.15), I(0.7), I(0.4), direction="left", color=ORANGE).name = f"plug {k + 1}"
add_text(s, X0, I(5.45), CONTENT_W, I(0.36),
         [[("플랫폼이 돌리고, ", {"role": "bridge"}), ("규칙은 우리가 쓴다.", {"role": "bridge", "color": ORANGE_TEXT})]],
         role="bridge").name = "click2 · bridge"
add_cite(s, "Example source — https://example.com/platform-guide")
drop_empty_placeholders(s)

# ── 3 · answers: frames and veils are new shapes on top; line arrows = correspondence
# start = cards → click1 = B → click2 = C → click3 = A · D dimmed → click4 = what does not change
s = content_slide(prs, "빨리 올리려면 관리형, 익숙하면 컨테이너")
CARD_Y, CARD_H = 1.3, 1.6
bottom = CARD_Y + CARD_H
cards = option_cards(add_group(s, "start · options"), CARD_Y, CARD_H, "body")
for step, (i, answer, why) in enumerate([(1, "컨테이너 플랫폼", "이미 쓰는 도구 그대로"),
                                         (2, "관리형 서비스", "가장 빨리 올리려면")], 1):
    g = add_group(s, f"click{step} · {'ABCD'[i]}")
    add_frame(g, cards[i]).name = f"frame {'ABCD'[i]}"
    cx = card_x(i) + CARD_W / 2
    add_arrow(g, I(cx), I(bottom + 0.03), I(cx), I(bottom + 0.45), color=ACCENT, weight=1.5).name = "answer arrow"
    add_text(g, I(card_x(i)), I(bottom + 0.5), I(CARD_W), I(0.6),
             [[(answer, {"role": "label", "color": ACCENT})], [(why, {"role": "body"})]],
             role="body", align=PP_ALIGN.CENTER, line_spacing=1.0).name = f"answer {'ABCD'[i]}"
g = add_group(s, "click3 · dimmed")
for i, why in [(0, "운영 인력이 따로 든다"), (3, "기한을 못 맞춘다")]:
    add_veil(g, cards[i]).name = f"veil {'ABCD'[i]}"
    add_text(g, I(card_x(i)), I(bottom + 0.5), I(CARD_W), I(0.3), [[(why, {"role": "body", "color": MUTED})]],
             role="body", align=PP_ALIGN.CENTER, line_spacing=1.0).name = f"reason {'ABCD'[i]}"
add_text(s, X0, I(4.75), CONTENT_W, I(0.36),
         [[("어느 쪽이든, ", {"role": "bridge"}), ("경보 규칙은 직접 정의한다.", {"role": "bridge", "color": ORANGE_TEXT})]],
         role="bridge").name = "click4 · invariant"
drop_empty_placeholders(s)

count_text(prs)
check_slides(prs)     # stops here on a violation — nothing is saved
os.makedirs("out", exist_ok=True)
prs.save("out/example_light.pptx")
print("saved out/example_light.pptx slides:", len(prs.slides))
